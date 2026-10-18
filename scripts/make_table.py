import sys, glob, os, re, collections
sys.path.insert(0, os.path.dirname(__file__))
import scriptfmt as S
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = "decoded/Chapter/Pack"
COLS = [("ja", None), ("en", "ue"), ("fr", "fr"), ("de", "de"), ("es", "es"), ("it", "it")]


def texts(path):
    if not os.path.exists(path):
        return {}
    out, seen = {}, {}
    for i, b in S.parse(S.read_s(path))[1]:
        if not i:
            continue
        seen[i] = seen.get(i, 0) + 1
        out[i if seen[i] == 1 else f"{i}#{seen[i]}"] = b   # repeated IDs in one file get a #N suffix
    return out


def clean(b):
    """-> (text, comments). Lines starting with // are developer comments, not shown in game."""
    lines = b.replace("\r\n", "\n").split("\n")
    text, com = [], []
    for ln in lines:
        if ln.startswith("//"):
            if ln.strip("/ "):
                com.append(ln[2:].strip())
        else:
            text.append(ln)
    return "\n".join(text).replace("\n<R>\n", "\n").strip("\n"), " | ".join(com)


def speakers(scene_txt):
    """id -> face id from last <SETFACE2> before <TEXT_EX>, plus reading order."""
    if not os.path.exists(scene_txt):
        return {}, []
    t = open(scene_txt, "rb").read().decode("utf-16", errors="replace")
    face, out, order = "", {}, []
    for m in re.finditer(r"<SETFACE2>\s*(\S+)|<TEXT_EX>\s*(\S+)", t):
        if m.group(1):
            face = m.group(1)
        else:
            out[m.group(2)] = face; order.append(m.group(2))
    return out, order


rows = []
bases = set()
for p in glob.glob(f"{ROOT}/*/*.s") + glob.glob(f"{ROOT}/*/*/*.s"):
    d, b = os.path.split(p)
    if os.path.basename(d) in ("ue", "fr", "de", "es", "it"):
        d = os.path.dirname(d)
    bases.add((d, b))
for d, base in sorted(bases):
    ja = f"{d}/{base}"; chap = os.path.basename(d); scene = base[:-2]
    data = {"ja": texts(ja)}
    for name, dirn in COLS[1:]:
        data[name] = texts(f"{d}/{dirn}/{base}")
    face, order = speakers(f"{d}/{scene}.txt")
    ids = list(dict.fromkeys(order + list(data["ja"]) + list(data["en"])))
    ids = [i for i in ids if any(i in data[n] for n in data)]
    for i in ids:
        cl = {n: clean(data[n].get(i, "")) for n, _ in COLS}
        notes = " || ".join(f"{n}: {cl[n][1]}" for n, _ in COLS if cl[n][1])
        rows.append([chap, scene, i, face.get(i, "")] + [cl[n][0] for n, _ in COLS] + ["", notes])

# ---- plain text files outside the chapters: whole file = one entry (chapter column = folder, ID = "-")
RAW_DIRS = ["2D", "EdoComment", "Profile"]


def raw_text(path):
    if not os.path.exists(path):
        return ""
    t = open(path, "rb").read().decode("utf-16").replace("\r\n", "\n")
    return t.strip("\n")


for d in RAW_DIRS:
    names = sorted(os.path.basename(x) for x in glob.glob(f"extracted/data/{d}/ue/*.txt"))
    for fn in names:
        cells = [raw_text(f"extracted/data/{d}/{dirn}/{fn}") if dirn else "" for name, dirn in COLS]
        rows.append([d, fn[:-4], "-", ""] + cells + ["", ""])

# keep an existing translation: the Russian column of the old table is re-applied by (chapter, scene, id)
OLD_CSV = "text/ru.csv" if os.path.exists("text/ru.csv") else "text/dialogue.csv"   # ru.csv is tracked in git
if os.path.exists(OLD_CSV):
    import csv as _csv
    _old = list(_csv.reader(open(OLD_CSV, encoding="utf-8-sig")))
    if _old and "Русский" in _old[0]:
        ri = _old[0].index("Русский")
        kept = {(x[0], x[1], x[2]): x[ri] for x in _old[1:] if len(x) > ri and x[ri].strip()}
        n_kept = 0
        for r in rows:
            v = kept.get((r[0], r[1], r[2]))
            if v:
                r[10] = v; n_kept += 1
        print(f"kept {n_kept} Russian cells from the previous table ({len(kept)} found)")

wb = Workbook(); ws = wb.active; ws.title = "Реплики"
hdr = ["Глава", "Сцена", "ID", "Лицо (face id)", "Японский", "English", "Français", "Deutsch", "Español", "Italiano", "Русский", "Комментарии в файле (//)"]
ws.append(hdr)
for r in rows:
    ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="444444")
for col, w in zip("ABCDEFGHIJKL", [8, 12, 8, 8, 34, 34, 28, 28, 28, 28, 34, 30]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "E2"; ws.auto_filter.ref = ws.dimensions
os.makedirs("text", exist_ok=True)
wb.save("text/dialogue.xlsx")
import csv
with open("text/dialogue.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f); w.writerow(hdr); w.writerows(rows)
print(len(rows), "rows;", len({r[0] for r in rows}), "chapters")
